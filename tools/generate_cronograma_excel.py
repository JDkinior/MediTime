import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_cronograma_workbook():
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------------------
    # Typography & Palette: Segoe UI (Clean, Modern Executive Look)
    # -------------------------------------------------------------------------
    font_name = "Segoe UI"
    
    f_title = Font(name=font_name, size=13, bold=True, color="0F172A")
    f_subtitle = Font(name=font_name, size=9, italic=True, color="475569")
    f_super_header = Font(name=font_name, size=11, bold=True, color="FFFFFF")
    f_month = Font(name=font_name, size=10, bold=True, color="1E293B")
    f_week = Font(name=font_name, size=9, bold=True, color="334155")
    f_header_col = Font(name=font_name, size=10, bold=True, color="1E293B")
    
    f_proj = Font(name=font_name, size=9, bold=True, color="1E293B")
    f_act = Font(name=font_name, size=9, bold=False, color="1E293B")
    f_act_num = Font(name=font_name, size=9, bold=True, color="2563EB")
    
    # Cell "X" fonts: crisp white bold on vibrant background
    f_cell_x_white = Font(name=font_name, size=10, bold=True, color="FFFFFF")
    f_cell_x_dark = Font(name=font_name, size=10, bold=True, color="0F172A")
    
    f_kpi_title = Font(name=font_name, size=9, bold=True, color="334155")
    f_kpi_val = Font(name=font_name, size=9, bold=False, color="0F172A")
    f_kpi_bold = Font(name=font_name, size=9, bold=True, color="0F172A")
    
    # Alignments
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)
    
    # Borders: refined slate boundaries and soft inner dividers
    border_dark = Side(border_style="medium", color="334155")
    border_thin = Side(border_style="thin", color="94A3B8")
    border_light = Side(border_style="thin", color="CBD5E1")
    
    box_grid = Border(left=border_light, right=border_light, top=border_light, bottom=border_light)
    box_header = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    box_super = Border(left=border_dark, right=border_dark, top=border_dark, bottom=border_dark)
    
    # Elegant Color Fills
    fill_super = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Deep Slate
    fill_col_hdr = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid") # Slate-100
    fill_proj = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")    # Slate-50
    
    # Harmonious Pastel Months (refined aesthetic)
    fill_julio = PatternFill(start_color="FDEBD0", end_color="FDEBD0", fill_type="solid")   # Warm peach
    fill_agosto = PatternFill(start_color="D4EFDF", end_color="D4EFDF", fill_type="solid")  # Mint sage
    fill_sept = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")    # Soft sky blue
    fill_oct = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")     # Clean gray
    fill_nov = PatternFill(start_color="E8DAEF", end_color="E8DAEF", fill_type="solid")     # Soft lavender
    
    # Status Fills (Modern, vivid, professional)
    fill_fin = PatternFill(start_color="16A34A", end_color="16A34A", fill_type="solid")     # Emerald green 600
    fill_curso = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")   # Sky blue 600
    fill_pend = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")    # Red 600
    
    fill_fin_light = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_curso_light = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
    fill_pend_light = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    # =========================================================================
    # SHEET 1: Cronograma 2026-2 (Corte 3) - FORMATO RENOVADO Y ALTA ESTÉTICA
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Cronograma 2026-2 (C3)"
    ws1.views.sheetView[0].showGridLines = True
    
    # Main Titles
    ws1["A2"] = "2. Cronograma de actividades"
    ws1["A2"].font = f_title
    ws1["A3"] = "Proyecto de Grado PGC (Noveno Semestre) — Vigencia 2026 | MediTime v2.32.0"
    ws1["A3"].font = f_subtitle
    
    # Row Heights
    ws1.row_dimensions[2].height = 22
    ws1.row_dimensions[3].height = 16
    ws1.row_dimensions[5].height = 24
    ws1.row_dimensions[6].height = 22
    ws1.row_dimensions[7].height = 18
    
    # Super Header: "2026-2 (Segundo Semestre)" across cols C to V (20 cols)
    ws1.merge_cells("C5:V5")
    super_cell = ws1["C5"]
    super_cell.value = "PERIODO ACADÉMICO 2026-2"
    super_cell.font = f_super_header
    super_cell.alignment = align_center
    super_cell.fill = fill_super
    for col in range(3, 23):
        ws1.cell(row=5, column=col).border = box_header
        
    # Month Headers
    months_data = [
        ("Julio", 3, 6, fill_julio),
        ("Agosto", 7, 10, fill_agosto),
        ("Septiembre", 11, 14, fill_sept),
        ("Octubre", 15, 18, fill_oct),
        ("Noviembre", 19, 22, fill_nov),
    ]
    
    for m_name, c_s, c_e, m_fill in months_data:
        ws1.merge_cells(start_row=6, start_column=c_s, end_row=6, end_column=c_e)
        c = ws1.cell(row=6, column=c_s, value=m_name)
        c.font = f_month
        c.alignment = align_center
        c.fill = m_fill
        for col_i in range(c_s, c_e + 1):
            ws1.cell(row=6, column=col_i).border = box_header
            
    # Week Numbers
    for c_i in range(3, 23):
        w_num = ((c_i - 3) % 4) + 1
        c = ws1.cell(row=7, column=c_i, value=w_num)
        c.font = f_week
        c.alignment = align_center
        m_idx = (c_i - 3) // 4
        c.fill = months_data[m_idx][3]
        c.border = box_header
        
    # Col A & B Headers
    ws1.merge_cells("A5:A7")
    c_proj_h = ws1["A5"]
    c_proj_h.value = "PROYECTO"
    c_proj_h.font = f_super_header
    c_proj_h.alignment = align_center
    c_proj_h.fill = fill_super
    for r in range(5, 8):
        ws1.cell(row=r, column=1).border = box_header
        
    ws1.merge_cells("B5:B7")
    c_act_h = ws1["B5"]
    c_act_h.value = "ACTIVIDAD"
    c_act_h.font = f_super_header
    c_act_h.alignment = align_center
    c_act_h.fill = fill_super
    for r in range(5, 8):
        ws1.cell(row=r, column=2).border = box_header
        
    # =========================================================================
    # REVISED ACTIVITIES LIST (With User Requirements):
    # - "Pruebas piloto" moved further out (October weeks 3-4 instead of Sept)
    # - "Jornadas de apropiación" completely replaced with:
    #   "Preparación de versión Release (APK) y diseño de instrumentos de evaluación de usabilidad (SUS)"
    # - Pre-pilot engineering testing in Sept weeks 3-4
    # =========================================================================
    activities_2026_2 = [
        # (Name, [(col_idx_0_to_19, status)])
        # Julio (0..3), Agosto (4..7), Sept (8..11), Oct (12..15), Nov (16..19)
        ("1. Optimización algorítmica de cálculo de dosis en Progreso y modularización", [(0, "FIN"), (1, "FIN")]),
        ("2. Implementación de Widgets interactivos de escritorio para Android", [(1, "FIN"), (2, "FIN")]),
        ("3. Digitalización de recetas con Visión Artificial (OCR) y validación clínica", [(2, "FIN"), (3, "FIN")]),
        ("4. Rediseño visual en tema oscuro, bordes accesibles y adaptación Edge-to-Edge", [(3, "FIN"), (4, "FIN")]),
        ("5. Desarrollo del Módulo Cuidador con gestión multi-paciente y tema morado pastel", [(4, "FIN"), (5, "FIN"), (6, "FIN")]),
        ("6. Desarrollo del Módulo Animales (Veterinaria) independiente con paleta verde crema", [(6, "FIN"), (7, "FIN")]),
        ("7. Implementación de Onboarding temático guiado en 5 pasos e iconos reactivos en menú", [(7, "FIN"), (8, "FIN")]),
        ("8. Construcción de suite de pruebas unitarias automáticas (flutter test)", [(8, "FIN"), (9, "FIN")]),
        ("9. Compilación formal de documentación técnica C3 (Manuales y Plan de Sostenibilidad)", [(9, "FIN"), (10, "FIN")]),
        ("10. Pruebas de compatibilidad en dispositivos físicos heterogéneos y optimización de batería", [(10, "CURSO"), (11, "CURSO")]),
        ("11. Preparación de versión Release (APK firmado) y diseño de instrumentos de evaluación (SUS)", [(12, "CURSO"), (13, "CURSO")]),
        ("12. Ejecución de pruebas piloto de usabilidad con usuarios y cuidadores en entorno real", [(14, "PEND"), (15, "PEND")]),
        ("13. Análisis de resultados de las pruebas piloto y sistematización de métricas de adherencia", [(16, "PEND"), (17, "PEND")]),
        ("14. Consolidación de informe final de transferencia, entrega de repositorio y sustentación", [(18, "PEND"), (19, "PEND")]),
    ]
    
    start_row = 8
    for idx, (act_name, weeks_data) in enumerate(activities_2026_2):
        curr_row = start_row + idx
        ws1.row_dimensions[curr_row].height = 25
        
        # Activity cell
        c_act = ws1.cell(row=curr_row, column=2, value=act_name)
        c_act.font = f_act
        c_act.alignment = align_left
        c_act.border = box_grid
        
        # Subtle zebra striping for activity rows
        if idx % 2 == 1:
            c_act.fill = fill_proj
            
        # Blank week cells with clean grid
        for w in range(20):
            c_w = ws1.cell(row=curr_row, column=3 + w)
            c_w.border = box_grid
            c_w.alignment = align_center
            if idx % 2 == 1:
                c_w.fill = fill_proj
                
        # Fill active weeks
        for w_idx, status in weeks_data:
            c_active = ws1.cell(row=curr_row, column=3 + w_idx, value="X")
            c_active.font = f_cell_x_white
            c_active.alignment = align_center
            c_active.border = box_header
            if status == "FIN":
                c_active.fill = fill_fin
            elif status == "CURSO":
                c_active.fill = fill_curso
            elif status == "PEND":
                c_active.fill = fill_pend
                
    end_row = start_row + len(activities_2026_2) - 1
    
    # Merge Project Name Column (Col A)
    ws1.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
    c_proj = ws1.cell(row=start_row, column=1, value="Aplicación móvil para la gestión de tratamientos médicos (MediTime)")
    c_proj.font = f_proj
    c_proj.alignment = align_center
    c_proj.fill = fill_col_hdr
    for r in range(start_row, end_row + 1):
        ws1.cell(row=r, column=1).border = box_header
        
    # =========================================================================
    # EXECUTIVE LEGEND & KPI SUMMARY CARD (Clean, Professional)
    # =========================================================================
    leg_start = end_row + 2
    
    # Legend Table Header
    ws1.cell(row=leg_start, column=2, value="ESTADO DE ACTIVIDADES").font = f_month
    ws1.cell(row=leg_start, column=2).alignment = align_center
    ws1.cell(row=leg_start, column=2).fill = fill_super
    ws1.cell(row=leg_start, column=2).font = f_super_header
    ws1.cell(row=leg_start, column=2).border = box_header
    
    ws1.cell(row=leg_start, column=3, value="ID").font = f_super_header
    ws1.cell(row=leg_start, column=3).alignment = align_center
    ws1.cell(row=leg_start, column=3).fill = fill_super
    ws1.cell(row=leg_start, column=3).border = box_header
    
    ws1.cell(row=leg_start, column=4, value="CANT.").font = f_super_header
    ws1.cell(row=leg_start, column=4).alignment = align_center
    ws1.cell(row=leg_start, column=4).fill = fill_super
    ws1.cell(row=leg_start, column=4).border = box_header
    
    ws1.merge_cells(start_row=leg_start, start_column=5, end_row=leg_start, end_column=6)
    c_porc = ws1.cell(row=leg_start, column=5, value="% PARTICIPACIÓN")
    c_porc.font = f_super_header
    c_porc.alignment = align_center
    c_porc.fill = fill_super
    ws1.cell(row=leg_start, column=5).border = box_header
    ws1.cell(row=leg_start, column=6).border = box_header
    
    legend_data = [
        ("Finalizado (Implementado y verificado)", "X", fill_fin, 9, "64.3%"),
        ("En Curso (En ejecución activa)", "X", fill_curso, 2, "14.3%"),
        ("Pendiente (Programado según cronograma)", "X", fill_pend, 3, "21.4%"),
    ]
    
    for l_idx, (l_text, l_symbol, l_fill, l_count, l_pct) in enumerate(legend_data):
        r_l = leg_start + 1 + l_idx
        ws1.row_dimensions[r_l].height = 20
        
        c_desc = ws1.cell(row=r_l, column=2, value=l_text)
        c_desc.font = f_kpi_val
        c_desc.alignment = align_left
        c_desc.border = box_grid
        
        c_sym = ws1.cell(row=r_l, column=3, value=l_symbol)
        c_sym.font = f_cell_x_white
        c_sym.alignment = align_center
        c_sym.fill = l_fill
        c_sym.border = box_header
        
        c_cnt = ws1.cell(row=r_l, column=4, value=l_count)
        c_cnt.font = f_kpi_bold
        c_cnt.alignment = align_center
        c_cnt.border = box_grid
        
        ws1.merge_cells(start_row=r_l, start_column=5, end_row=r_l, end_column=6)
        c_p = ws1.cell(row=r_l, column=5, value=l_pct)
        c_p.font = f_kpi_bold
        c_p.alignment = align_center
        c_p.border = box_grid
        ws1.cell(row=r_l, column=6).border = box_grid
        
    # Total row
    r_total = leg_start + 4
    ws1.row_dimensions[r_total].height = 21
    c_tot_lbl = ws1.cell(row=r_total, column=2, value="TOTAL GENERAL Y PROGRESO")
    c_tot_lbl.font = f_kpi_bold
    c_tot_lbl.alignment = align_left
    c_tot_lbl.fill = fill_col_hdr
    c_tot_lbl.border = box_header
    
    c_tot_sym = ws1.cell(row=r_total, column=3, value="")
    c_tot_sym.fill = fill_col_hdr
    c_tot_sym.border = box_header
    
    c_tot_cnt = ws1.cell(row=r_total, column=4, value=14)
    c_tot_cnt.font = f_kpi_bold
    c_tot_cnt.alignment = align_center
    c_tot_cnt.fill = fill_col_hdr
    c_tot_cnt.border = box_header
    
    ws1.merge_cells(start_row=r_total, start_column=5, end_row=r_total, end_column=6)
    c_tot_p = ws1.cell(row=r_total, column=5, value="100.0% (Avance: 78.6%)")
    c_tot_p.font = f_kpi_bold
    c_tot_p.alignment = align_center
    c_tot_p.fill = fill_col_hdr
    c_tot_p.border = box_header
    ws1.cell(row=r_total, column=6).border = box_header

    # Column Widths Optimization
    ws1.column_dimensions["A"].width = 18
    ws1.column_dimensions["B"].width = 64
    for c_idx in range(3, 23):
        col_letter = get_column_letter(c_idx)
        ws1.column_dimensions[col_letter].width = 4.2

    # =========================================================================
    # SHEET 2: Cronograma Anual Consolidado 2026 (Febrero a Noviembre)
    # =========================================================================
    ws2 = wb.create_sheet(title="Cronograma Anual 2026")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2["A2"] = "2. Cronograma Maestro de Desarrollo e Innovación (Vigencia 2026)"
    ws2["A2"].font = f_title
    ws2["A3"] = "Visión Macro-Evolutiva Anual (Febrero – Noviembre 2026) | MediTime v2.32.0"
    ws2["A3"].font = f_subtitle
    
    ws2.row_dimensions[5].height = 24
    ws2.row_dimensions[6].height = 22
    ws2.row_dimensions[7].height = 18
    
    ws2.merge_cells("C5:V5")
    s2 = ws2["C5"]
    s2.value = "AÑO ACADÉMICO Y TECNOLÓGICO 2026"
    s2.font = f_super_header
    s2.alignment = align_center
    s2.fill = fill_super
    for col in range(3, 23):
        ws2.cell(row=5, column=col).border = box_header
        
    annual_months = [
        ("Febrero", 3, 4, fill_julio),
        ("Marzo", 5, 6, fill_agosto),
        ("Abril", 7, 8, fill_sept),
        ("Mayo", 9, 10, fill_oct),
        ("Junio", 11, 12, fill_nov),
        ("Julio", 13, 14, fill_julio),
        ("Agosto", 15, 16, fill_agosto),
        ("Septiembre", 17, 18, fill_sept),
        ("Octubre", 19, 20, fill_oct),
        ("Noviembre", 21, 22, fill_nov),
    ]
    
    for m_name, c_s, c_e, m_f in annual_months:
        ws2.merge_cells(start_row=6, start_column=c_s, end_row=6, end_column=c_e)
        c = ws2.cell(row=6, column=c_s, value=m_name)
        c.font = f_month
        c.alignment = align_center
        c.fill = m_f
        for col_i in range(c_s, c_e + 1):
            ws2.cell(row=6, column=col_i).border = box_header
            
    for c_i in range(3, 23):
        q_label = "Q1" if (c_i % 2 != 0) else "Q2"
        c = ws2.cell(row=7, column=c_i, value=q_label)
        c.font = f_week
        c.alignment = align_center
        c.border = box_header
        
    ws2.merge_cells("A5:A7")
    ws2["A5"] = "PROYECTO"
    ws2["A5"].font = f_super_header
    ws2["A5"].alignment = align_center
    ws2["A5"].fill = fill_super
    for r in range(5, 8): ws2.cell(row=r, column=1).border = box_header
    
    ws2.merge_cells("B5:B7")
    ws2["B5"] = "ACTIVIDAD"
    ws2["B5"].font = f_super_header
    ws2["B5"].alignment = align_center
    ws2["B5"].fill = fill_super
    for r in range(5, 8): ws2.cell(row=r, column=2).border = box_header
    
    annual_activities = [
        ("1. Integración de ShowcaseView para tutorial guiado interactivo y onboarding inicial", [(0, "FIN"), (1, "FIN"), (2, "FIN")]),
        ("2. Actualización de servicios base de Google y refactorización de dependencias", [(2, "FIN"), (3, "FIN")]),
        ("3. Análisis de costos y migración de almacenamiento a Cloudinary con auto-foto Google", [(6, "FIN"), (7, "FIN")]),
        ("4. Implementación del Localizador de Farmacias con geolocalización y mapa interactivo", [(6, "FIN"), (7, "FIN")]),
        ("5. Integración del Asistente Virtual Bilingüe Midi con Groq API (IA LLaMA y Gemini)", [(6, "FIN"), (7, "FIN"), (8, "FIN")]),
        ("6. Rediseño visual Serene Health y adaptación a temas moderno / oscuro", [(8, "FIN"), (9, "FIN")]),
        ("7. Implementación de adherencia terapéutica real y dashboard interactivo de progreso", [(8, "FIN"), (9, "FIN")]),
        ("8. Generación y exportación de reportes clínicos certificados en formato PDF", [(8, "FIN"), (9, "FIN")]),
        ("9. Desarrollo de Widgets interactivos de pantalla de inicio para Android", [(10, "FIN"), (11, "FIN")]),
        ("10. Digitalización de recetas con Visión Artificial (OCR) y validación de interacciones", [(10, "FIN"), (11, "FIN")]),
        ("11. Desarrollo del Módulo de Cuidador (Caregiver Mode) con paleta morado pastel", [(12, "FIN"), (13, "FIN")]),
        ("12. Desarrollo del Módulo Animales y Veterinaria independiente con paleta verde crema", [(13, "FIN"), (14, "FIN")]),
        ("13. Onboarding guiado en 5 pasos, adaptación Edge-to-Edge y menú reactivo dinámico", [(14, "FIN"), (15, "FIN")]),
        ("14. Suite de pruebas unitarias automáticas (flutter test) y refactorización", [(14, "FIN"), (15, "FIN")]),
        ("15. Elaboración y compilación de documentos institucionales C3 (Manuales y Sostenibilidad)", [(15, "FIN"), (16, "FIN")]),
        ("16. Pruebas de compatibilidad en dispositivos físicos y optimización de batería", [(16, "CURSO"), (17, "CURSO")]),
        ("17. Preparación de versión Release (APK) y diseño de instrumentos de evaluación (SUS)", [(17, "CURSO"), (18, "CURSO")]),
        ("18. Ejecución de pruebas piloto de usabilidad con usuarios y cuidadores en entorno real", [(18, "PEND"), (19, "PEND")]),
        ("19. Consolidación de informe final de entrega, repositorios y sustentación de grado", [(19, "PEND"), (19, "PEND")]),
    ]
    
    start_r2 = 8
    for idx, (act_name, q_data) in enumerate(annual_activities):
        c_r = start_r2 + idx
        ws2.row_dimensions[c_r].height = 24
        
        c_act = ws2.cell(row=c_r, column=2, value=act_name)
        c_act.font = f_act
        c_act.alignment = align_left
        c_act.border = box_grid
        if idx % 2 == 1:
            c_act.fill = fill_proj
            
        for q in range(20):
            c_q = ws2.cell(row=c_r, column=3 + q)
            c_q.border = box_grid
            c_q.alignment = align_center
            if idx % 2 == 1:
                c_q.fill = fill_proj
                
        for q_idx, status in q_data:
            c_w = ws2.cell(row=c_r, column=3 + q_idx, value="X")
            c_w.font = f_cell_x_white
            c_w.alignment = align_center
            c_w.border = box_header
            if status == "FIN":
                c_w.fill = fill_fin
            elif status == "CURSO":
                c_w.fill = fill_curso
            elif status == "PEND":
                c_w.fill = fill_pend
                
    end_r2 = start_r2 + len(annual_activities) - 1
    ws2.merge_cells(start_row=start_r2, start_column=1, end_row=end_r2, end_column=1)
    c_proj2 = ws2.cell(row=start_r2, column=1, value="Aplicación móvil para la gestión de tratamientos médicos (MediTime)")
    c_proj2.font = f_proj
    c_proj2.alignment = align_center
    c_proj2.fill = fill_col_hdr
    for r in range(start_r2, end_r2 + 1):
        ws2.cell(row=r, column=1).border = box_header
        
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 66
    for c_idx in range(3, 23):
        col_letter = get_column_letter(c_idx)
        ws2.column_dimensions[col_letter].width = 4.4

    # =========================================================================
    # SHEET 3: Cierre y Evolución del Cronograma Original 2025-1
    # =========================================================================
    ws3 = wb.create_sheet(title="Cierre Cronograma 2025-1")
    ws3.views.sheetView[0].showGridLines = True
    
    ws3["A2"] = "2. Estado de Cierre y Trazabilidad del Cronograma Original (2025-1)"
    ws3["A2"].font = f_title
    ws3["A3"] = "Auditoría de Cumplimiento Técnico al 100% sobre las 8 Actividades Iniciales"
    ws3["A3"].font = f_subtitle
    
    ws3.row_dimensions[5].height = 24
    
    headers_cierre = [
        ("N°", 6),
        ("Actividad Original (2025-1)", 38),
        ("Estado en Imagen Antigua", 22),
        ("Estado Real Actual", 20),
        ("Commit Git Clave", 22),
        ("Resolución Técnica y Evolución Posterior", 54)
    ]
    
    for c_idx, (h_text, h_w) in enumerate(headers_cierre, start=1):
        cell = ws3.cell(row=5, column=c_idx, value=h_text)
        cell.font = f_super_header
        cell.alignment = align_center
        cell.fill = fill_super
        cell.border = box_header
        col_letter = get_column_letter(c_idx)
        ws3.column_dimensions[col_letter].width = h_w
        
    cierre_data = [
        (
            1,
            "Implementación del tutorial interactivo para nuevos usuarios",
            "Finalizado (Verde)",
            "FINALIZADO (100%)",
            "ca8e51b / cd18109",
            "Consolidado con ShowcaseView interactivo y ampliado a onboarding de 5 pasos con diferenciación temática."
        ),
        (
            2,
            "Diseño de flujo de navegación del tutorial interactivo",
            "Finalizado (Verde)",
            "FINALIZADO (100%)",
            "ca8e51b / 139c1d3",
            "Flujo de navegación interactivo integrado a pantallas principales (Home, Progreso, Calendario y Opciones)."
        ),
        (
            3,
            "Integración del tutorial con la interfaz existente de la aplicación",
            "En Curso (Azul)",
            "FINALIZADO (100%)",
            "ca8e51b / 139c1d3",
            "Completamente integrado sin interferir con la experiencia del usuario y con opción de repetición en Ayuda."
        ),
        (
            4,
            "Análisis de cambios en políticas y costos de Firebase",
            "En Curso (Azul)",
            "FINALIZADO (100%)",
            "af224ac / CLOUDINARY...",
            "Se diagnosticó la saturación de cuotas en Firebase Storage, lo que motivó la migración estratégica a Cloudinary."
        ),
        (
            5,
            "Optimización del uso de Firebase (lecturas, escrituras y almacenamiento)",
            "En Curso (Azul)",
            "FINALIZADO (100%)",
            "da51878 / dc3ff1d / af224ac",
            "Arquitectura Limpia con repositorios Firestore desacoplados, cache local y migración de multimedia a Cloudinary."
        ),
        (
            6,
            "Corrección de errores en notificaciones",
            "Pendiente (Rojo)",
            "FINALIZADO (100%)",
            "a2aaa1f / 850fc9d / 617cb60",
            "Motor de hardware con Android AlarmManager exacto, RebootBroadcastReceiver y ActionBroadcastReceiver en 2do plano."
        ),
        (
            7,
            "Actualización del manual de usuario con el nuevo tutorial",
            "Pendiente (Rojo)",
            "FINALIZADO (100%)",
            "documentos/C3_Manual...",
            "Ampliación a dossier institucional formal C3 (Manual de Usuario de 14 capítulos y Manual Técnico de Arquitectura)."
        ),
        (
            8,
            "Pruebas de rendimiento después de las optimizaciones",
            "Pendiente (Rojo)",
            "FINALIZADO (100%)",
            "97a4440 / flutter test",
            "Optimización algorítmica de cálculo de dosis, lazy loading de tratamientos indefinidos y suite de tests aprobada al 100%."
        ),
    ]
    
    for row_idx, row_vals in enumerate(cierre_data, start=6):
        ws3.row_dimensions[row_idx].height = 26
        for col_idx, val in enumerate(row_vals, start=1):
            c = ws3.cell(row=row_idx, column=col_idx, value=val)
            c.font = f_act
            c.border = box_grid
            if col_idx in (1, 4):
                c.alignment = align_center
                if col_idx == 4:
                    c.fill = fill_fin_light
                    c.font = Font(name=font_name, size=9, bold=True, color="16A34A")
            elif col_idx == 3:
                c.alignment = align_center
            else:
                c.alignment = align_left

    output_path = "documentos/Cronograma_de_Actividades_MediTime_2026.xlsx"
    actualizado_path = "documentos/Cronograma_de_Actividades_MediTime_Actualizado.xlsx"
    wb.save(output_path)
    wb.save(actualizado_path)
    print(f"Refined Excel successfully saved at: {output_path} and {actualizado_path}")

if __name__ == "__main__":
    build_cronograma_workbook()
